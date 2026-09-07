import pytest
from app.parser import parse_document


def test_parse_markdown():
    data = "# 标题\n\nMySQL 连接超时排查步骤。\n\n- 检查 max_connections\n".encode("utf-8")
    text = parse_document("runbook.md", data)
    assert "MySQL 连接超时" in text
    assert "max_connections" in text


def test_parse_txt():
    text = parse_document("faq.txt", "常见错误码 1040 表示连接数已满。".encode("utf-8"))
    assert "1040" in text


def test_parse_docx(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("巡检标准操作流程第一段。")
    doc.add_paragraph("第二段：检查磁盘使用率。")
    buf = tmp_path / "sop.docx"
    doc.save(buf)
    text = parse_document("sop.docx", buf.read_bytes())
    assert "巡检标准操作流程" in text
    assert "磁盘使用率" in text


def test_parse_pdf(tmp_path):
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buf = tmp_path / "blank.pdf"
    with open(buf, "wb") as f:
        writer.write(f)
    # 空白页应解析为空字符串而不是报错
    assert parse_document("blank.pdf", buf.read_bytes()) == ""


def test_unsupported_type():
    with pytest.raises(ValueError, match="不支持的文件类型"):
        parse_document("a.exe", b"MZ")
